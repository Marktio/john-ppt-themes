#!/usr/bin/env python3
"""
tools/make_looks.py -- builds the starter look zips from looks_src/.

  looks_src/looks.json          the list of looks (names, colors, fonts, layouts)
  looks_src/engine_template.js  the shared drawing code (layout archetypes)

For each look it: writes engine.js + manifest.json, runs a REAL test deck
(mix mode, single-layout mode, every theme, a no-photo slide, a long title,
a 9-bullet slide), checks that no text sits outside the safe margin, renders a
thumbnail with LibreOffice (if installed), and zips it into looks/<kind>/<id>.zip
with a matching <id>.png preview next to it.

Needs: Python 3, Node, and `npm install pptxgenjs` somewhere Node can find it
(set NODE_PATH to that node_modules folder). Optional: python-pptx (margin
check), Pillow (placeholder photos), LibreOffice + pdftoppm (thumbnails).

Usage:
    python tools/make_looks.py                 # build all
    python tools/make_looks.py --only slate-pro
    python tools/make_looks.py --no-render     # skip thumbnails (keeps old .png)
Safe to repeat: it overwrites only the zips/pngs of the looks it builds.
"""
from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
SRC = HERE / "looks_src"
KIND_DIR = {"standard": REPO / "looks" / "standard", "static": REPO / "looks" / "static"}

SAMPLE_OUTLINE = """## A Calm Look At Deck Design
- Subtitle line for the opening slide
## Why layouts matter
- A photo-led slide shows the picture first
- Text stays short and easy to scan
- Colors come from the theme, never hard-coded
## This is a deliberately long title to prove that sixty plus characters still fit the box
- One
- Two
- Three
- Four
- Five
- Six
- Seven
- Eight
- Nine
## ★ Part Two
## Another content slide with a photo
- Bullet one
- Bullet two
- Bullet three
## Thank you
- Questions welcome
"""


def run(cmd, **kw):
    return subprocess.run(cmd, check=True, capture_output=True, text=True, **kw)


def make_images(d: Path):
    d.mkdir(parents=True, exist_ok=True)
    try:
        from PIL import Image
    except ImportError:
        return  # photo-less test only
    for n, col in ((2, (120, 150, 190)), (5, (190, 150, 120))):
        # slide 3 and 1 deliberately have NO photo
        from PIL import ImageDraw
        img = Image.new("RGB", (900, 900), col)
        draw = ImageDraw.Draw(img)
        for y in range(900):  # soft top-to-bottom shade, stands in for a photo
            shade = int(70 * y / 900)
            draw.line([(0, y), (900, y)], fill=(max(col[0] - shade, 0), max(col[1] - shade, 0), max(col[2] - shade, 0)))
        img.save(d / f"{n}.png")


def check_margins(pptx: Path, label: str):
    """Text boxes must stay inside 0.3in..9.7in (x) and 0.3in..5.3in (y). Full-bleed pictures are exempt."""
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        print(f"    (margin check skipped for {label}: python-pptx not installed)")
        return
    prs = Presentation(str(pptx))
    bad = []
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            x, y, w, h = (Emu(v).inches for v in (sh.left, sh.top, sh.width, sh.height))
            # the slide-number footer sits at y=5.0..5.3 by design
            if x < 0.299 or y < 0.299 or x + w > 9.701 or y + h > 5.301:
                bad.append((i, sh.text_frame.text[:30], round(x, 2), round(y, 2), round(x + w, 2), round(y + h, 2)))
    if bad:
        raise SystemExit(f"MARGIN FAIL {label}: {bad}")


def check_no_dropped_text(pptx: Path, label: str):
    """Every bullet in the sample outline must appear on its slide (nothing silently cut)."""
    try:
        from pptx import Presentation
    except ImportError:
        return
    expected, cur = [], None
    for ln in SAMPLE_OUTLINE.splitlines():
        if ln.startswith("## "):
            cur = []
            expected.append(cur)
        elif ln.startswith("- ") and cur is not None:
            cur.append(ln[2:].strip())
    prs = Presentation(str(pptx))
    for i, (slide, want) in enumerate(zip(prs.slides, expected), 1):
        have = " | ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
        missing = [b for b in want if b not in have]
        if missing:
            raise SystemExit(f"DROPPED TEXT {label} slide {i}: {missing}")


def render_thumb(pptx: Path, out_png: Path) -> bool:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    ppm = shutil.which("pdftoppm")
    if not (soffice and ppm):
        return False
    work = pptx.parent
    run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(work), str(pptx)], timeout=180)
    pdf = work / (pptx.stem + ".pdf")
    run([ppm, "-r", "110", "-png", "-f", "2", "-l", "2", "-scale-to-x", "1100", "-scale-to-y", "-1", str(pdf), str(work / "thumb")])
    made = sorted(work.glob("thumb*.png"))
    if not made:
        return False
    shutil.copyfile(made[0], out_png)
    return True


def build_one(look: dict, template: str, render: bool):
    name = look["name"]
    print(f"== {name} ({look['kind']})")
    tmp = Path(tempfile.mkdtemp(prefix="look-"))
    pack = tmp / name
    pack.mkdir()
    cfg = {k: look[k] for k in ("name", "themes", "defaultTheme", "deco", "layouts", "types", "content")}
    (pack / "engine.js").write_text(template.replace("__LOOK__", json.dumps(cfg, indent=2)), encoding="utf-8")
    manifest = {
        "name": name,
        "description": look["description"],
        "engine": "engine.js",
        "layoutCount": len(look["layouts"]),
        "themes": list(look["themes"].keys()),
        "defaultTheme": look["defaultTheme"],
        "supportsMultiImage": False,
        "variants": [],
    }
    (pack / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")

    outline = tmp / "outline.md"
    outline.write_text(SAMPLE_OUTLINE, encoding="utf-8")
    imgs = tmp / "imgs"
    make_images(imgs)

    runs = [("mix", None)] + [(str(n), None) for n in (1, len(look["layouts"]))] + [("mix", t) for t in look["themes"]]
    first_mix = None
    for layout, theme in runs:
        out = tmp / f"sample-{layout}-{theme or 'default'}.pptx"
        cmd = ["node", str(pack / "engine.js"), str(outline), str(imgs), str(out), f"--layout={layout}"]
        if theme:
            cmd.append(f"--theme={theme}")
        run(cmd)
        if not out.is_file() or out.stat().st_size < 5000:
            raise SystemExit(f"{name}: no real .pptx produced for layout={layout} theme={theme}")
        check_margins(out, f"{name}/{layout}/{theme}")
        check_no_dropped_text(out, f"{name}/{layout}/{theme}")
        if layout == "mix" and theme is None:
            first_mix = out
        print(f"   ok  layout={layout:<4} theme={theme or 'default'}")
    # unknown theme id must not throw
    run(["node", str(pack / "engine.js"), str(outline), str(imgs), str(tmp / "x.pptx"), "--theme=does-not-exist"])
    print("   ok  unknown theme falls back to default")

    dest_dir = KIND_DIR[look["kind"]]
    dest_dir.mkdir(parents=True, exist_ok=True)
    thumb_made = False
    if render and first_mix:
        thumb_made = render_thumb(first_mix, pack / "thumbnail.png")
        print("   ok  thumbnail rendered" if thumb_made else "   --  thumbnail skipped (LibreOffice/pdftoppm not found)")
    zpath = dest_dir / f"{look['id']}.zip"
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(pack.iterdir()):
            z.write(f, f"{name}/{f.name}")
    if thumb_made:
        shutil.copyfile(pack / "thumbnail.png", dest_dir / f"{look['id']}.png")
    print(f"   wrote {zpath.relative_to(REPO)}")
    shutil.rmtree(tmp, ignore_errors=True)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--only", help="build just this look id")
    ap.add_argument("--no-render", action="store_true")
    a = ap.parse_args()
    if not shutil.which("node"):
        sys.exit("Node.js not found. Install it (winget install OpenJS.NodeJS.LTS) and run npm install pptxgenjs.")
    template = (SRC / "engine_template.js").read_text(encoding="utf-8")
    looks = json.loads((SRC / "looks.json").read_text(encoding="utf-8"))["looks"]
    if a.only:
        looks = [l for l in looks if l["id"] == a.only]
        if not looks:
            sys.exit(f"No look with id {a.only}")
    for look in looks:
        build_one(look, template, not a.no_render)
    print("Done. Now run update-manifest.bat (or python build_manifest.py).")


if __name__ == "__main__":
    main()
